"""Excel structure and typed import/export helpers for the Ananas system."""

from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SHEET_NAME = "Plantations Ananas"
COORDINATE_HEADER = "Coordonnées Géographiques"
HEADERS = [
    "N",
    "Nom et prénoms",
    "Sexe",
    "Commune",
    "Arrondissement",
    "Village/Hameau",
    COORDINATE_HEADER,
    None,  # La cellule est comprise dans l'en-tête fusionné G1:H1 de la pièce jointe.
    "Superficie déclarée (ha)",
    "Superficie trackée (ha)",
    "Type de friche",
    "Type d’espèce végétale",
    "Type de sol",
    "Précédents culturaux",
    "Décision de l'équipe de validation",
    "Décision ATDA7",
    "Sup attribuée / confirmée (ha)",
]

WIDTHS = [8, 28, 10, 18, 20, 22, 16, 16, 20, 20, 22, 24, 20, 24, 28, 20, 24]
TEXT_FIELDS = {
    1: "nom_prenoms",
    2: "sexe",
    3: "commune",
    4: "arrondissement",
    5: "village_hameau",
    10: "type_friche",
    11: "type_espece_vegetale",
    12: "type_sol",
    13: "precedents_culturaux",
    14: "decision_equipe_validation",
    15: "decision_atda7",
}
NUMERIC_FIELDS = {
    6: "coord_x",
    7: "coord_y",
    8: "superficie_declaree",
    9: "superficie_trackee",
    16: "superficie_confirmee",
}


def _decimal_value(value: Any, *, integer: bool = False) -> Decimal | int | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise ValueError("une valeur booléenne n'est pas un nombre valide")
    try:
        parsed = Decimal(str(value).strip().replace(" ", "").replace(",", "."))
    except (InvalidOperation, ValueError):
        raise ValueError("une valeur numérique est attendue")
    if not parsed.is_finite():
        raise ValueError("le nombre doit être fini")
    if integer:
        if parsed != parsed.to_integral_value():
            raise ValueError("un nombre entier est attendu")
        return int(parsed)
    if parsed < 0:
        raise ValueError("la valeur ne peut pas être négative")
    return parsed


def _text_value(value: Any, header: str) -> str | None:
    if value is None:
        return None
    if not isinstance(value, str):
        raise ValueError(f"une valeur texte est attendue dans « {header} »")
    cleaned = value.strip()
    return cleaned or None


def _validate_headers(ws) -> None:
    actual = [ws.cell(1, column).value for column in range(1, len(HEADERS) + 1)]
    normalized = [value.strip() if isinstance(value, str) else value for value in actual]
    if normalized != HEADERS:
        raise ValueError(
            "Les en-têtes de la première ligne ne correspondent pas au modèle Ananas. "
            "La cellule G1 doit contenir « Coordonnées Géographiques » et être fusionnée avec H1."
        )
    if any(ws.cell(1, column).value not in (None, "") for column in range(len(HEADERS) + 1, ws.max_column + 1)):
        raise ValueError("Le classeur contient des colonnes supplémentaires après « Sup attribuée / confirmée (ha) ».")
    if "G1:H1" not in {str(rng) for rng in ws.merged_cells.ranges}:
        raise ValueError("L'en-tête « Coordonnées Géographiques » doit couvrir les colonnes X et Y (G1:H1).")
    subheaders = (ws.cell(2, 7).value, ws.cell(2, 8).value)
    if subheaders != ("X", "Y"):
        raise ValueError("La deuxième ligne doit contenir « X » en G et « Y » en H sous l'en-tête fusionné.")
    if any(ws.cell(2, column).value not in (None, "") for column in (*range(1, 7), *range(9, len(HEADERS) + 1))):
        raise ValueError("La deuxième ligne du modèle doit contenir uniquement les coordonnées X et Y.")
    if any(ws.cell(2, column).value not in (None, "") for column in range(len(HEADERS) + 1, ws.max_column + 1)):
        raise ValueError("Le classeur contient des en-têtes supplémentaires sur la deuxième ligne.")


def parse_import_workbook(content: bytes) -> tuple[list[dict], list[dict]]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=False, data_only=True)
    except Exception as exc:
        raise ValueError("Le fichier Excel est illisible ou endommagé.") from exc

    ws = workbook.active
    if ws.max_row < 1:
        raise ValueError("Le classeur ne contient aucune ligne d'en-tête.")
    _validate_headers(ws)

    valid_rows: list[dict] = []
    row_errors: list[dict] = []
    for row_number in range(3, ws.max_row + 1):
        values = [ws.cell(row_number, column).value for column in range(1, len(HEADERS) + 1)]
        if all(value is None or (isinstance(value, str) and not value.strip()) for value in values):
            continue

        errors = []
        if any(ws.cell(row_number, column).value not in (None, "") for column in range(len(HEADERS) + 1, ws.max_column + 1)):
            errors.append("des données dépassent la dernière colonne du modèle")
        payload: dict[str, Any] = {}
        try:
            numero = _decimal_value(values[0], integer=True)
            if numero is None:
                raise ValueError("« N » est obligatoire")
            payload["numero"] = numero
        except ValueError as exc:
            errors.append(str(exc))

        for index, field in TEXT_FIELDS.items():
            try:
                payload[field] = _text_value(values[index], HEADERS[index])
            except ValueError as exc:
                errors.append(str(exc))

        for index, field in NUMERIC_FIELDS.items():
            try:
                parsed = _decimal_value(values[index])
                if field in ("coord_x", "coord_y") and parsed is not None:
                    if parsed != parsed.to_integral_value():
                        raise ValueError("des coordonnées entières sont attendues")
                payload[field] = parsed
            except ValueError as exc:
                errors.append(f"« {HEADERS[index]} » : {exc}")

        if errors:
            row_errors.append({"ligne": row_number, "erreurs": errors})
        else:
            valid_rows.append(payload)

    workbook.close()
    return valid_rows, row_errors


def build_template_workbook() -> Workbook:
    workbook = Workbook()
    ws = workbook.active
    ws.title = SHEET_NAME
    for index, value in enumerate(HEADERS, start=1):
        if value is not None:
            ws.cell(1, index, value)
    ws.merge_cells("G1:H1")
    for index in (*range(1, 7), *range(9, len(HEADERS) + 1)):
        letter = get_column_letter(index)
        ws.merge_cells(f"{letter}1:{letter}2")
    ws.cell(2, 7, "X")
    ws.cell(2, 8, "Y")
    header_fill = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")
    for index, width in enumerate(WIDTHS, start=1):
        for row in (1, 2):
            cell = ws.cell(row, index)
            cell.fill = header_fill
            cell.font = Font(bold=True, color="000000")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[1].height = 44
    ws.row_dimensions[2].height = 18
    ws.freeze_panes = "A3"
    return workbook


def _row_from_item(item: dict) -> list[Any]:
    return [
        item.get("numero"), item.get("nom_prenoms"), item.get("sexe"),
        item.get("commune"), item.get("arrondissement"), item.get("village_hameau"),
        item.get("coord_x"), item.get("coord_y"), item.get("superficie_declaree"),
        item.get("superficie_trackee"), item.get("type_friche"),
        item.get("type_espece_vegetale"), item.get("type_sol"),
        item.get("precedents_culturaux"), item.get("decision_equipe_validation"),
        item.get("decision_atda7"), item.get("superficie_confirmee"),
    ]


def build_data_workbook(sheets: list[tuple[str, list[dict]]]) -> Workbook:
    workbook = Workbook()
    first = True
    for title, rows in sheets:
        ws = workbook.active if first else workbook.create_sheet()
        first = False
        ws.title = title[:31]
        for index, value in enumerate(HEADERS, start=1):
            if value is not None:
                ws.cell(1, index, value)
        ws.merge_cells("G1:H1")
        for index in (*range(1, 7), *range(9, len(HEADERS) + 1)):
            letter = get_column_letter(index)
            ws.merge_cells(f"{letter}1:{letter}2")
        ws.cell(2, 7, "X")
        ws.cell(2, 8, "Y")
        header_fill = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")
        for index, width in enumerate(WIDTHS, start=1):
            for row in (1, 2):
                cell = ws.cell(row, index)
                cell.fill = header_fill
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.row_dimensions[1].height = 44
        ws.row_dimensions[2].height = 18
        for item in rows:
            ws.append(_row_from_item(item))
        ws.freeze_panes = "A3"
    return workbook
