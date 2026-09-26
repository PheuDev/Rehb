import csv
import io
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response, StreamingResponse
from sqlalchemy.orm import Session

from app import crud, excel_utils, schemas
from app.completion import missing_fields
from app.dependencies import get_db, require_active, require_admin, require_chef_or_admin
from app.models import User
from app.crud import SORTABLE_COLUMNS, SUP_CLASSES

router = APIRouter(prefix="/api/rehabilitations", tags=["Réhabilitations"])

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _workbook_response(workbook, filename: str) -> StreamingResponse:
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def _scope(db: Session, user: User) -> Optional[list[str]]:
    """Périmètre de visibilité des fiches pour l'utilisateur courant.

    - ``None`` → administrateur : toutes les fiches ;
    - liste    → fiches des brigades affectées à l'équipe de l'utilisateur.
    """
    return crud.get_visible_brigade_names(db, user)


# ---------------------------------------------------------------------------
# Liste paginée avec recherche / filtres / tri
# ---------------------------------------------------------------------------
@router.get("", response_model=schemas.RehabilitationListResponse)
def list_rehabilitations(
    q: Optional[str] = Query(None, description="Recherche plein texte"),
    departement: Optional[str] = None,
    commune: Optional[str] = None,
    arrondissement: Optional[str] = None,
    village: Optional[str] = None,
    brigade_name: Optional[str] = None,
    annee: Optional[int] = None,
    sup_class: Optional[str] = Query(None, description="Classe de superficie"),
    page: int = Query(1, ge=1),
    limit: int = Query(10, ge=1, le=200),
    sortBy: str = Query("created_at"),
    sortOrder: str = Query("desc", pattern="^(asc|desc)$"),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    if sup_class is not None and sup_class not in SUP_CLASSES:
        raise HTTPException(status_code=400, detail="Classe de superficie invalide.")
    if sortBy not in SORTABLE_COLUMNS:
        raise HTTPException(status_code=400, detail="Champ de tri invalide.")

    items, pagination = crud.get_list(
        db,
        q=q,
        departement=departement,
        commune=commune,
        arrondissement=arrondissement,
        village=village,
        brigade_name=brigade_name,
        annee=annee,
        sup_class=sup_class,
        page=page,
        limit=limit,
        sort_by=sortBy,
        sort_order=sortOrder,
        brigade_names=_scope(db, current_user),
    )
    return {"items": items, "pagination": pagination}


@router.get("/incomplete", response_model=schemas.RehabilitationIncompleteListResponse)
def list_incomplete_rehabilitations(
    q: Optional[str] = Query(None, description="Recherche plein texte"),
    departement: Optional[str] = None,
    commune: Optional[str] = None,
    arrondissement: Optional[str] = None,
    village: Optional[str] = None,
    brigade_name: Optional[str] = None,
    annee: Optional[int] = None,
    sup_class: Optional[str] = Query(None, description="Classe de superficie"),
    page: int = Query(1, ge=1),
    limit: int = Query(10, ge=1, le=200),
    sortBy: str = Query("created_at"),
    sortOrder: str = Query("desc", pattern="^(asc|desc)$"),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    """Retourne les fiches qui possèdent au moins une donnée à renseigner."""
    if sup_class is not None and sup_class not in SUP_CLASSES:
        raise HTTPException(status_code=400, detail="Classe de superficie invalide.")
    if sortBy not in SORTABLE_COLUMNS:
        raise HTTPException(status_code=400, detail="Champ de tri invalide.")

    items, pagination = crud.get_incomplete_list(
        db, q, departement, commune, arrondissement, village, brigade_name,
        annee, sup_class, page, limit, sortBy, sortOrder,
        brigade_names=_scope(db, current_user),
    )
    response_items = []
    for item in items:
        data = schemas.RehabilitationOut.model_validate(item).model_dump()
        data["missingFields"] = missing_fields(item)
        response_items.append(data)
    return {"items": response_items, "pagination": pagination}


# ---------------------------------------------------------------------------
# Valeurs distinctes pour les listes déroulantes de filtres
# ---------------------------------------------------------------------------
@router.get("/filters", response_model=schemas.FiltersResponse)
def get_filters(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    """Valeurs distinctes pour les listes déroulantes (limitées au périmètre)."""
    return crud.get_filters(db, brigade_names=_scope(db, current_user))


# ---------------------------------------------------------------------------
# Liste des brigades disponibles (regroupées par nom)
# ---------------------------------------------------------------------------
# NB : route statique, à déclarer avant la route dynamique "/{rehab_id}".
@router.get("/brigades", response_model=list[schemas.BrigadeOut])
def list_brigades(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    """Brigades distinctes regroupées par nom (limitées au périmètre)."""
    return crud.get_brigades(db, brigade_names=_scope(db, current_user))


@router.get("/brigades-detail", response_model=schemas.BrigadeDetailListResponse)
def list_brigades_detail(
    q: Optional[str] = Query(None, description="Recherche par nom, responsable, commune, village…"),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    """Brigades avec toutes les informations agrégées (responsable, superficie, communes…)."""
    return crud.get_brigades_detail(db, q=q, brigade_names=_scope(db, current_user))


@router.get("/brigades-detail/export-excel")
def export_brigades_excel(
    q: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    """Exporte la liste des brigades (agrégée) au format Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    data = crud.get_brigades_detail(db, q=q, brigade_names=_scope(db, current_user))
    items = data["items"]

    HEADER_FILL = PatternFill(start_color="2D6846", end_color="2D6846", fill_type="solid")
    HEADER_FONT = Font(color="FFFFFF", bold=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Brigades"

    headers = [
        "Nom de la brigade", "Chef de brigade", "Téléphone",
        "Nb fiches", "Superficie totale (ha)",
        "Nb communes", "Nb villages", "Nb départements",
        "Communes", "Départements", "Années actives",
    ]
    col_widths = [30, 28, 20, 12, 22, 14, 14, 18, 40, 40, 25]

    ws.append(headers)
    for idx, (_, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=1, column=idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    for b in items:
        ws.append([
            b["brigade_name"],
            b["manager_name"] or "",
            b["manager_phone"] or "",
            b["fiches"],
            b["superficie_totale"],
            b["nb_communes"],
            b["nb_villages"],
            b["nb_departements"],
            ", ".join(b["communes"]),
            ", ".join(b["departements"]),
            ", ".join(str(a) for a in b["annees"]),
        ])

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 1)}"
    return _workbook_response(wb, "brigades.xlsx")


@router.get("/departements/export-excel")
def export_departements_excel(
    q: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    """Exporte la liste des départements (agrégée) au format Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    data = crud.get_departements(db, q=q, brigade_names=_scope(db, current_user))
    items = data["items"]

    HEADER_FILL = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")
    HEADER_FONT = Font(color="FFFFFF", bold=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Départements"

    headers = [
        "Département", "Nb fiches", "Superficie totale (ha)",
        "Nb communes", "Nb arrondissements", "Nb villages",
        "Communes", "Brigades", "Années actives",
    ]
    col_widths = [25, 12, 22, 14, 20, 14, 45, 45, 25]

    ws.append(headers)
    for idx, (_, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=1, column=idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    for d in items:
        ws.append([
            d["departement"],
            d["fiches"],
            d["superficie_totale"],
            d["nb_communes"],
            d["nb_arrondissements"],
            d["nb_villages"],
            ", ".join(d["communes"]),
            ", ".join(d["brigades"]),
            ", ".join(str(a) for a in d["annees"]),
        ])

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 1)}"
    return _workbook_response(wb, "departements.xlsx")


@router.get("/producteurs/export-excel")
def export_producteurs_excel(
    q: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    """Exporte la liste des producteurs (agrégée) au format Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    data = crud.get_producers(db, q=q, brigade_names=_scope(db, current_user))
    items = data["items"]

    HEADER_FILL = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
    HEADER_FONT = Font(color="FFFFFF", bold=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Producteurs"

    headers = [
        "Nom du producteur", "Téléphone",
        "Nb fiches", "Superficie totale (ha)",
        "Communes", "Villages", "Brigades",
    ]
    col_widths = [35, 20, 12, 22, 40, 40, 40]

    ws.append(headers)
    for idx, (_, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=1, column=idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    for p in items:
        ws.append([
            p["producer_name"],
            p["producer_phone"] or "",
            p["fiches"],
            p["superficie_totale"],
            ", ".join(p["communes"]),
            ", ".join(p["villages"]),
            ", ".join(p["brigades"]),
        ])

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 1)}"
    return _workbook_response(wb, "producteurs.xlsx")


# ---------------------------------------------------------------------------
# Audit superficies — échantillonnage aléatoire
# ---------------------------------------------------------------------------

def _saved_audit_summary(row) -> dict:
    user = row.created_by
    return {
        "id": row.id,
        "title": row.title,
        "created_at": row.created_at,
        "created_by_username": user.username if user else None,
        "created_by_full_name": user.full_name if user else None,
        "nb_fiches_echantillon": row.nb_fiches_echantillon,
        "superficie_echantillon": float(row.superficie_echantillon) if row.superficie_echantillon is not None else None,
        "pourcentage_couverture": float(row.pourcentage_couverture) if row.pourcentage_couverture is not None else None,
        "brigade_filter": row.brigade_filter,
    }


def _build_audit_excel_workbook(result: dict):
    """Construit le classeur Excel du plan d'audit (échantillon + hors échantillon + synthèse)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    HEADER_FILL = PatternFill(start_color="2D6846", end_color="2D6846", fill_type="solid")
    HEADER_FONT = Font(color="FFFFFF", bold=True)
    BRIGADE_FILL = PatternFill(start_color="DCE0DC", end_color="DCE0DC", fill_type="solid")
    BRIGADE_FONT = Font(bold=True)

    COL_HEADERS = ["N°PDA", "Brigade", "Classe de superficie", "Departement", "Commune",
                   "Arrondissement", "Village", "Producteur", "Superficie (ha)", "Annee"]

    def _write_headers(ws):
        ws.append(COL_HEADERS)
        for i, _ in enumerate(COL_HEADERS, 1):
            c = ws.cell(row=1, column=i)
            c.fill = HEADER_FILL
            c.font = HEADER_FONT
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = 22
        ws.row_dimensions[1].height = 28
        ws.freeze_panes = "A2"

    def _write_brigade_block(ws, brigade_name, fiches, row_num):
        title_cell = ws.cell(row=row_num, column=1, value=f"  Brigade : {brigade_name}  ({len(fiches)} fiche(s))")
        title_cell.fill = BRIGADE_FILL
        title_cell.font = BRIGADE_FONT
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=len(COL_HEADERS))
        row_num += 1
        for r in fiches:
            sup = float(r.superficie_rehabilitee) if r.superficie_rehabilitee else None
            ws.append([
                r.pda_number,
                r.brigade_name,
                crud._audit_class(sup) if sup else None,
                r.departement, r.commune, r.arrondissement, r.village,
                r.producer_name, sup, r.annee_rehabilitation,
            ])
            row_num += 1
        return row_num

    wb = Workbook()

    ws_a = wb.active
    ws_a.title = "A - Echantillon audit"
    _write_headers(ws_a)

    brigades_a: dict[str, list] = {}
    for r in result["fiches"]:
        key = r.brigade_name or "— Sans brigade —"
        brigades_a.setdefault(key, []).append(r)

    row_a = 2
    for brigade_name in sorted(brigades_a.keys()):
        row_a = _write_brigade_block(ws_a, brigade_name, brigades_a[brigade_name], row_a)
    ws_a.append([])
    ws_a.append(["", "TOTAL ECHANTILLON", "", "", "", "", "", "",
                 result["superficie_echantillon"],
                 f"{result['pourcentage_couverture']} % de la superficie totale"])

    ws_b = wb.create_sheet("B - Hors echantillon")
    _write_headers(ws_b)

    brigades_b: dict[str, list] = {}
    for r in result["fiches_hors_echantillon"]:
        key = r.brigade_name or "— Sans brigade —"
        brigades_b.setdefault(key, []).append(r)

    row_b = 2
    for brigade_name in sorted(brigades_b.keys()):
        row_b = _write_brigade_block(ws_b, brigade_name, brigades_b[brigade_name], row_b)

    ws_c = wb.create_sheet("C - Synthese par brigade")
    headers_c = ["Brigade", "Total fiches", "Fiches echantillon", "Couverture fiches (%)",
                 "Superficie brigade (ha)", "Superficie echantillon (ha)", "Couverture superficie (%)"]
    ws_c.append(headers_c)
    for i, _ in enumerate(headers_c, 1):
        c = ws_c.cell(row=1, column=i)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        ws_c.column_dimensions[get_column_letter(i)].width = 26
    ws_c.freeze_panes = "A2"
    for b in result["par_brigade"]:
        ws_c.append([b["brigade"], b["total_fiches"], b["fiches_echantillon"],
                     f"{b['pourcentage_fiches']} %",
                     b["superficie_brigade"], b["superficie_echantillon"],
                     f"{b['pourcentage_superficie']} %"])

    return wb


@router.get("/audit-sample", response_model=schemas.AuditSampleResponse)
def get_audit_sample(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    """Génère un plan d'échantillonnage aléatoire pour l'audit des superficies."""
    result = crud.get_audit_sample(db, brigade_names=_scope(db, current_user))

    # Enrichir chaque fiche avec sa classe d'audit
    fiches_out = []
    for r in result["fiches"]:
        sup = float(r.superficie_rehabilitee) if r.superficie_rehabilitee else None
        fiche_data = schemas.AuditFicheOut.model_validate(r)
        fiche_data.audit_classe = crud._audit_class(sup) if sup else None
        fiches_out.append(fiche_data)


    # Enrichir les fiches hors echantillon avec leur classe d'audit
    hors_echantillon_out = []
    for r in result["fiches_hors_echantillon"]:
        sup = float(r.superficie_rehabilitee) if r.superficie_rehabilitee else None
        fiche_data = schemas.AuditFicheOut.model_validate(r)
        fiche_data.audit_classe = crud._audit_class(sup) if sup else None
        hors_echantillon_out.append(fiche_data)

    return {
        "fiches": fiches_out,
        "fiches_hors_echantillon": hors_echantillon_out,
        "total_fiches": result["total_fiches"],
        "superficie_echantillon": result["superficie_echantillon"],
        "superficie_totale": result["superficie_totale"],
        "pourcentage_couverture": result["pourcentage_couverture"],
        "par_classe": result["par_classe"],
        "par_brigade": result["par_brigade"],
    }


@router.get("/audit-sample/export-excel")
def export_audit_excel(
    brigades: Optional[str] = Query(None, description="Noms de brigades séparés par virgule"),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    """Exporte le plan d'audit en Excel : Feuille A (echantillon) + Feuille B (hors echantillon)."""
    result = crud.get_audit_sample(db, brigade_names=_scope(db, current_user))

    brigade_filter = [b.strip() for b in brigades.split(",")] if brigades else None
    if brigade_filter:
        result = dict(result)
        result["fiches"] = [r for r in result["fiches"] if (r.brigade_name or "— Sans brigade —") in brigade_filter]
        result["fiches_hors_echantillon"] = [r for r in result["fiches_hors_echantillon"] if (r.brigade_name or "— Sans brigade —") in brigade_filter]
        result["par_brigade"] = [b for b in result["par_brigade"] if b["brigade"] in brigade_filter]

    wb = _build_audit_excel_workbook(result)
    return _workbook_response(wb, "plan_audit_superficies.xlsx")


@router.post("/audit-suggestions", response_model=schemas.SavedAuditSuggestionDetail, status_code=201)
def save_audit_suggestion(
    payload: schemas.SavedAuditSuggestionCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    """Enregistre une suggestion d'échantillon d'audit pour consultation ultérieure."""
    from datetime import datetime

    title = (payload.title or "").strip()
    if not title:
        title = f"Suggestion du {datetime.now().strftime('%d/%m/%Y %H:%M')}"

    snapshot = payload.snapshot.model_dump(mode="json")
    row = crud.create_saved_audit_suggestion(
        db,
        user=current_user,
        title=title,
        snapshot=snapshot,
        brigade_filter=payload.brigade_filter or None,
    )
    summary = _saved_audit_summary(row)
    return {**summary, "snapshot": payload.snapshot}


@router.get("/audit-suggestions", response_model=schemas.SavedAuditSuggestionListResponse)
def list_audit_suggestions(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    rows = crud.list_saved_audit_suggestions(db)
    items = [_saved_audit_summary(r) for r in rows]
    return {"items": items, "total": len(items)}


@router.get("/audit-suggestions/{suggestion_id}", response_model=schemas.SavedAuditSuggestionDetail)
def get_audit_suggestion(
    suggestion_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    row = crud.get_saved_audit_suggestion(db, suggestion_id)
    if not row:
        raise HTTPException(status_code=404, detail="Suggestion introuvable.")
    return {**_saved_audit_summary(row), "snapshot": row.snapshot}


@router.get("/audit-suggestions/{suggestion_id}/export-excel")
def export_saved_audit_excel(
    suggestion_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    row = crud.get_saved_audit_suggestion(db, suggestion_id)
    if not row:
        raise HTTPException(status_code=404, detail="Suggestion introuvable.")
    result = crud.audit_snapshot_to_export_result(row.snapshot)
    wb = _build_audit_excel_workbook(result)
    safe_title = "".join(c if c.isalnum() or c in "._-" else "_" for c in row.title)[:60]
    return _workbook_response(wb, f"audit_{safe_title or suggestion_id}.xlsx")


@router.delete("/audit-suggestions/{suggestion_id}", status_code=204)
def delete_audit_suggestion(
    suggestion_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    if not crud.delete_saved_audit_suggestion(db, suggestion_id):
        raise HTTPException(status_code=404, detail="Suggestion introuvable.")
    return Response(status_code=204)


# ---------------------------------------------------------------------------
# Liste des départements
# ---------------------------------------------------------------------------
@router.get("/departements", response_model=schemas.DepartementListResponse)
def list_departements(
    q: Optional[str] = Query(None, description="Recherche par nom de département"),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    """Départements distincts avec leurs agrégats (fiches, superficie, communes…)."""
    return crud.get_departements(db, q=q, brigade_names=_scope(db, current_user))


# ---------------------------------------------------------------------------
# Liste des producteurs (regroupés par nom)
# ---------------------------------------------------------------------------
@router.get("/producteurs", response_model=schemas.ProducerListResponse)
def list_producers(
    q: Optional[str] = Query(None, description="Recherche par nom, téléphone, commune ou village"),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    """Producteurs distincts avec leurs informations agrégées (fiches, superficie, communes…)."""
    return crud.get_producers(db, q=q, brigade_names=_scope(db, current_user))


# ---------------------------------------------------------------------------
# Statistiques globales et agrégats
# ---------------------------------------------------------------------------
@router.get("/stats", response_model=schemas.StatsResponse)
def get_stats(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    """Statistiques (périmètre complet pour un admin, équipe pour les autres)."""
    return crud.get_stats(db, brigade_names=_scope(db, current_user))
# ---------------------------------------------------------------------------
# Export CSV (respecte les filtres actifs)
# ---------------------------------------------------------------------------
@router.get("/export")
def export_csv(
    q: Optional[str] = None,
    departement: Optional[str] = None,
    commune: Optional[str] = None,
    arrondissement: Optional[str] = None,
    village: Optional[str] = None,
    brigade_name: Optional[str] = None,
    annee: Optional[int] = None,
    sup_class: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    rows = crud.get_all_for_export(
        db, q, departement, commune, arrondissement, village, brigade_name, annee, sup_class,
        brigade_names=_scope(db, current_user),
    )

    buffer = io.StringIO()
    buffer.write("\ufeff")  # BOM UTF-8 pour Excel
    writer = csv.writer(buffer, delimiter=";")

    headers = [
        "N°PDA", "Département", "Commune", "Arrondissement", "Village",
        "Année", "Superficie réhabilitée (ha)", "Classe de superficie",
        "Nom de la brigade", "Responsable brigade", "Téléphone responsable",
        "Producteur bénéficiaire", "Téléphone producteur",
        "Désherbage (ha)", "Opérateur désherbage", "Téléphone opérateur désherbage",
        "Éclaircie (ha)", "Opérateur éclaircie", "Téléphone opérateur éclaircie",
        "Élagage (ha)", "Opérateur élagage", "Téléphone opérateur élagage",
        "Débardage (ha)", "Opérateur débardage", "Téléphone opérateur débardage",
        "Observations",
    ]
    writer.writerow(headers)

    for r in rows:
        writer.writerow([
            r.pda_number, r.departement, r.commune, r.arrondissement, r.village,
            r.annee_rehabilitation, r.superficie_rehabilitee, r.sup_class,
            r.brigade_name or "", r.brigade_manager_name or "", r.brigade_manager_phone or "",
            r.producer_name or "", r.producer_phone or "",
            r.desherbage_superficie or "", r.desherbage_operateur_nom or "", r.desherbage_operateur_phone or "",
            r.eclaircie_superficie or "", r.eclaircie_operateur_nom or "", r.eclaircie_operateur_phone or "",
            r.elagage_superficie or "", r.elagage_operateur_nom or "", r.elagage_operateur_phone or "",
            r.debardage_superficie or "", r.debardage_operateur_nom or "", r.debardage_operateur_phone or "",
            r.observations or "",
        ])

    buffer.seek(0)
    return StreamingResponse(
        iter([buffer.getvalue()]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=rehabilitations.csv"},
    )


# ---------------------------------------------------------------------------
# Modèle Excel à télécharger (mêmes colonnes que le fichier source)
# ---------------------------------------------------------------------------
@router.get("/import-template")
def download_import_template(current_user: User = Depends(require_active)):
    workbook = excel_utils.build_template_workbook()
    return _workbook_response(workbook, "modele_import_rehabilitations.xlsx")


# ---------------------------------------------------------------------------
# Export Excel (respecte les filtres actifs)
# ---------------------------------------------------------------------------
@router.get("/export-excel")
def export_excel(
    q: Optional[str] = None,
    departement: Optional[str] = None,
    commune: Optional[str] = None,
    arrondissement: Optional[str] = None,
    village: Optional[str] = None,
    brigade_name: Optional[str] = None,
    annee: Optional[int] = None,
    sup_class: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    rows = crud.get_all_for_export(
        db, q, departement, commune, arrondissement, village, brigade_name, annee, sup_class,
        brigade_names=_scope(db, current_user),
    )
    workbook = excel_utils.build_export_workbook(rows)
    return _workbook_response(workbook, "rehabilitations.xlsx")


@router.get("/export-incomplete-excel")
def export_incomplete_excel(
    q: Optional[str] = None,
    departement: Optional[str] = None,
    commune: Optional[str] = None,
    arrondissement: Optional[str] = None,
    village: Optional[str] = None,
    brigade_name: Optional[str] = None,
    annee: Optional[int] = None,
    sup_class: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    rows = crud.get_all_incomplete_for_export(
        db, q, departement, commune, arrondissement, village, brigade_name, annee, sup_class,
        brigade_names=_scope(db, current_user),
    )
    workbook = excel_utils.build_incomplete_export_workbook(rows)
    return _workbook_response(workbook, "fiches_a_completer.xlsx")


# ---------------------------------------------------------------------------
# Import Excel (structure identique au fichier source / au modèle téléchargeable)
# ---------------------------------------------------------------------------
@router.post("/import-excel", response_model=schemas.ImportResult)
async def import_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_chef_or_admin),
):
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Le fichier doit être au format Excel (.xlsx).")

    content = await file.read()
    try:
        valid_payloads, row_errors, incomplete_count = excel_utils.parse_import_workbook(content)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:  # pragma: no cover - fichier corrompu, mauvais format, etc.
        raise HTTPException(status_code=400, detail=f"Impossible de lire le fichier Excel : {exc}") from exc

    imported = []
    if valid_payloads:
        try:
            imported = crud.bulk_create(db, valid_payloads)
        except Exception as exc:  # pragma: no cover
            db.rollback()
            raise HTTPException(status_code=400, detail=f"Erreur lors de l'import : {exc}") from exc

    return {
        "total": len(valid_payloads) + len(row_errors),
        "importees": len(imported),
        "aCompleter": incomplete_count,
        "erreurs": row_errors,
    }


# ---------------------------------------------------------------------------
# Détail
# ---------------------------------------------------------------------------
@router.get("/{rehab_id}", response_model=schemas.RehabilitationOut)
def get_rehabilitation(
    rehab_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    obj = crud.get_by_id(db, rehab_id)
    if not obj:
        raise HTTPException(status_code=404, detail="Fiche de réhabilitation introuvable.")
    # Un non-admin ne peut consulter que les fiches de son périmètre
    if not crud.is_fiche_visible(db, current_user, obj.brigade_name):
        raise HTTPException(
            status_code=403,
            detail="Cette fiche n'appartient pas aux brigades de votre équipe.",
        )
    return obj


# ---------------------------------------------------------------------------
# Création
# ---------------------------------------------------------------------------
@router.post("", response_model=schemas.RehabilitationOut, status_code=201)
def create_rehabilitation(
    payload: schemas.RehabilitationCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    try:
        obj = crud.create(db, payload)
        # Lie automatiquement la fiche à son auteur (Phase 2)
        obj.author_id = current_user.id
        db.commit()
        db.refresh(obj)
        return obj
    except Exception as exc:  # pragma: no cover - erreurs DB génériques
        db.rollback()
        raise HTTPException(status_code=400, detail=f"Erreur lors de la création : {exc}") from exc


# ---------------------------------------------------------------------------
# Mise à jour
# ---------------------------------------------------------------------------
@router.put("/{rehab_id}", response_model=schemas.RehabilitationOut)
def update_rehabilitation(
    rehab_id: int,
    payload: schemas.RehabilitationUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active),
):
    obj = crud.get_by_id(db, rehab_id)
    if not obj:
        raise HTTPException(status_code=404, detail="Fiche de réhabilitation introuvable.")
    try:
        return crud.update(db, obj, payload)
    except Exception as exc:  # pragma: no cover
        db.rollback()
        raise HTTPException(status_code=400, detail=f"Erreur lors de la mise à jour : {exc}") from exc


# ---------------------------------------------------------------------------
# Suppression
# ---------------------------------------------------------------------------
# NB : la route statique "/all" (vidage) DOIT être déclarée AVANT la route
# dynamique "/{rehab_id}" : dans FastAPI, le routage se fait dans l'ordre de
# déclaration. Sinon "all" tente d'être converti en int et renvoie un 422.
@router.delete("/all")
def clear_rehabilitations(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    deleted = crud.clear_all(db)
    return {"deleted": deleted, "message": "Toutes les fiches ont été supprimées."}


@router.delete("/{rehab_id}")
def delete_rehabilitation(
    rehab_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_chef_or_admin),
):
    obj = crud.get_by_id(db, rehab_id)
    if not obj:
        raise HTTPException(status_code=404, detail="Fiche de réhabilitation introuvable.")
    crud.delete(db, obj)
    return {"message": "Fiche supprimée avec succès."}


@router.post("/bulk-delete")
def delete_rehabilitations(
    payload: schemas.BulkDeleteRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_chef_or_admin),
):
    ids = sorted(set(payload.ids))
    if not ids:
        raise HTTPException(status_code=400, detail="Aucune fiche sélectionnée.")

    deleted = crud.delete_many(db, ids)
    if deleted == 0:
        raise HTTPException(status_code=404, detail="Aucune fiche sélectionnée n’a été trouvée.")

    return {"deleted": deleted, "message": f"{deleted} fiche(s) supprimée(s) avec succès."}
